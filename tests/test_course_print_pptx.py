"""Course print (.docx) -> PowerPoint conversion.

Builds a tiny synthetic course print in the hand-made format (Title, Module
H1, "1 of 2 | Title" H2, layer H3, screenshots, "Audio Transcript:" label,
quiz-info line, answer table with an X) so no real course files are needed.
"""

import io
import os
import tempfile

from docx import Document
from docx.shared import Inches
from PIL import Image
from pptx import Presentation

from app.automation.course_print import docx_to_pptx as d2p


def _png(color="red"):
    buf = io.BytesIO()
    Image.new("RGB", (160, 90), color).save(buf, "PNG")
    buf.seek(0)
    return buf


def _bold(doc, text):
    doc.add_paragraph().add_run(text).bold = True


def _make_print(path, with_images=True):
    doc = Document()
    doc.add_paragraph("Sample Course", style="Title")
    doc.add_paragraph("Welcome", style="Heading 1")          # H1 with content -> screen
    if with_images:
        doc.add_picture(_png("blue"), width=Inches(6))
    _bold(doc, "Audio Transcript:")
    doc.add_paragraph("Welcome to the course.")
    doc.add_paragraph("Module 1: Basics", style="Heading 1")  # empty H1 -> divider
    doc.add_paragraph("1 of 2 | First Screen", style="Heading 2")
    if with_images:
        doc.add_picture(_png("red"), width=Inches(6))
        doc.add_picture(_png("green"), width=Inches(6))      # second state
    _bold(doc, "Audio Transcript:")
    doc.add_paragraph("Narration for screen one.")
    doc.add_paragraph("A bullet point", style="List Paragraph")
    doc.add_paragraph("Learn More (Slide Layer)", style="Heading 3")
    if with_images:
        doc.add_picture(_png("gray"), width=Inches(6))
    doc.add_paragraph("Layer text.")
    doc.add_paragraph("5 o 2 | Knowledge Check", style="Heading 2")   # typo seen in real prints
    doc.add_paragraph().add_run(" (Pick One, 10 points, unlimited attempts permitted)").italic = True
    _bold(doc, "Audio Transcript:")
    doc.add_paragraph("Which is right?")
    t = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Correct", "Choice"), (" X", "The right one"), ("", "The wrong one")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    doc.save(path)


def _convert(layout="notes", with_images=True):
    d = tempfile.mkdtemp()
    src = os.path.join(d, "print.docx")
    _make_print(src, with_images)
    summary = d2p.convert(src, layout=layout)
    return summary, Presentation(summary["out"])


def test_parse_structure():
    d = tempfile.mkdtemp()
    src = os.path.join(d, "p.docx")
    _make_print(src)
    cp = d2p.parse_course_print(src)
    assert cp.title == "Sample Course"
    assert [(n.level, n.heading) for n in cp.nodes] == [
        (1, "Welcome"), (1, "Module 1: Basics"), (2, "1 of 2 | First Screen"),
        (3, "Learn More (Slide Layer)"), (2, "5 o 2 | Knowledge Check")]
    first = cp.nodes[2]
    assert len(first.images) == 2
    assert [b.kind for b in first.blocks] == ["label", "para", "bullet"]
    quiz = cp.nodes[4]
    assert quiz.blocks[0].kind == "info"
    assert quiz.blocks[-1].kind == "table" and quiz.blocks[-1].answer_key


def test_slide_plan_and_deck():
    s, prs = _convert()
    # title + Welcome + divider + 2 views of screen 1 + layer + quiz (no shot -> text)
    assert s["slides"] == 7 and s["sections"] == 1 and s["images"] == 4 and s["text_only"] == 1
    assert prs.slide_width == d2p.SLIDE_W
    specs = d2p.plan_slides(d2p.parse_course_print(
        os.path.join(os.path.dirname(s["out"]), "print.docx")))
    kinds = [x.kind for x in specs]
    assert kinds == ["title", "screen", "section", "screen", "screen", "screen", "screen"]
    assert specs[2].eyebrow == "Module 1" and specs[2].title == "Basics"
    assert specs[3].pill == "1 of 2" and specs[3].tag == "View 1 of 2"
    assert specs[5].title == "Learn More" and specs[5].tag == "Layer"
    assert specs[6].pill == "5 of 2" and specs[6].title == "Knowledge Check"


def test_notes_carry_transcript_and_answer_key():
    _, prs = _convert()
    notes = [sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
             for sl in prs.slides]
    assert "Narration for screen one." in notes[3]
    assert "• A bullet point" in notes[3]
    assert "previous slide" in notes[4]                     # second view of same screen
    quiz = notes[6]
    assert "Pick One" in quiz
    assert "The right one  [Correct]" in quiz
    assert "The wrong one  [Correct]" not in quiz


def test_side_layout_puts_text_on_slide():
    _, prs = _convert(layout="side")
    texts = " ".join(sh.text_frame.text for sh in prs.slides[3].shapes if sh.has_text_frame)
    assert "Narration for screen one." in texts


def test_text_only_print_still_converts():
    s, prs = _convert(with_images=False)
    assert s["images"] == 0 and s["text_only"] == s["screen_slides"]
    texts = " ".join(sh.text_frame.text for sh in prs.slides[3].shapes if sh.has_text_frame)
    assert "Narration for screen one." in texts


def test_bad_accent_rejected():
    try:
        d2p._rgb("navy")
    except ValueError:
        return
    raise AssertionError("expected ValueError")
