"""
Microsoft Defender Attack Simulation Training — phish results report.

Two phases, matching the manual process this replaces
(examples/excel reports/ms-attack-sim_phase-*_process.txt):

Phase 1  raw results .csv  ->  "<Mon YYYY> <Ev|Sp> Phish Results.xlsx"
         Every user lands on exactly one sheet, first match wins (fail-first —
         a user who clicked AND reported is a failure because they clicked):

             Delivery Failed   FailedToDeliverEmail_TimeStamp set  (not counted)
             Fail              compromised / any click-or-open / replied
             forwards          forwarded (pending review)
             Pass - Report     Phishing Reported On set
             Pass - Delete     MessageDeleted_TimeStamp set
             Pass - Ignore     everyone left

         Plus "<Mon YYYY> <Ev|Sp> forwards.csv" with only UserName, UserMail,
         MessageForwarded_TimeStamp, Phishing Reported On, for the reviewer who
         checks whether each forward went to the approved backup reporter mailbox.

Phase 2  phase-1 workbook + forward review results + previous final workbook
         -> the same workbook, finished:
             forwarded to the approved mailbox  -> Pass - Report
             forwarded anywhere else             -> Fail
         then the Fail sheet is split by consecutive-failure streak using the
         previous simulation's final workbook:
             prev Fail x N           -> Fail x N+1
             prev Pass - * / absent  -> Fail x 1
             prev Delivery Failed    -> Fail x 1, flagged for a manual look back
         and the summary block for the results email is produced.

No Qt in here — the tab is app/ui/data_analysis_tab.py.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

# ---------------------------------------------------------------------------
# Sheet names / columns
# ---------------------------------------------------------------------------

SHEET_REPORT = "Pass - Report"
SHEET_DELETE = "Pass - Delete"
SHEET_IGNORE = "Pass - Ignore"
SHEET_FORWARDS = "forwards"
SHEET_FAIL = "Fail"
SHEET_DELIVERY_FAILED = "Delivery Failed"

# Tab order of the phase-1 workbook (as in ms attack sim filter 2.png)
PHASE1_SHEET_ORDER = [SHEET_REPORT, SHEET_DELETE, SHEET_IGNORE,
                      SHEET_FORWARDS, SHEET_FAIL, SHEET_DELIVERY_FAILED]

COL_NAME = "UserName"
COL_MAIL = "UserMail"
COL_COMPROMISED = "Compromised"
COL_DELIVERY_FAILED = "FailedToDeliverEmail_TimeStamp"
COL_REPLIED = "MessageReplied_TimeStamp"
COL_FORWARDED = "MessageForwarded_TimeStamp"
COL_REPORTED = "Phishing Reported On"
COL_DELETED = "MessageDeleted_TimeStamp"

FORWARDS_CSV_COLUMNS = [COL_NAME, COL_MAIL, COL_FORWARDED, COL_REPORTED]

# *_TimeStamp columns that are a user acting on the payload -> Fail.
_FAIL_PREFIXES = ("EmailLinkClicked", "AttachmentOpened", "CredSupplied")

# *_TimeStamp columns that are NOT a user acting on the payload.
# Any *_TimeStamp column in neither list is "unknown" (Microsoft added a
# technique or renamed something); the tab asks whether it counts as a fail.
# Called without an answer, unknown columns count as fail — errs towards
# Fail, never Pass.
_NON_INTERACTION_PREFIXES = (
    "SuccessfullyDeliveredEmail", "FailedToDeliverEmail",
    "MessageRead", "MessageDeleted", "MessageReplied", "MessageForwarded",
    "PositiveReinforcementMessage", "JustSimulationMessage",
    "TrainingAssignmentMessage",
)

_EMPTY = {"", "--"}
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_FAIL_X_RE = re.compile(r"^\s*fail\s*x\s*(\d+)\s*$", re.IGNORECASE)


def fail_x_sheet(n: int) -> str:
    return f"Fail x {n}"


def has_value(v) -> bool:
    """True if a cell holds real data (Defender writes "--" for "never happened")."""
    if v is None:
        return False
    return str(v).strip() not in _EMPTY


def _norm_email(v) -> str:
    return str(v).strip().lower() if has_value(v) else ""


# ---------------------------------------------------------------------------
# Table model
# ---------------------------------------------------------------------------

@dataclass
class Table:
    """A header row plus data rows (lists aligned to the header)."""
    headers: list[str]
    rows: list[list] = field(default_factory=list)

    def __post_init__(self):
        self._idx = {h.strip().lower(): i for i, h in enumerate(self.headers) if h}

    def col(self, name: str) -> int | None:
        return self._idx.get(name.strip().lower())

    def get(self, row: list, name: str):
        i = self.col(name)
        if i is None or i >= len(row):
            return None
        return row[i]

    def empty_like(self) -> "Table":
        return Table(list(self.headers))


# ---------------------------------------------------------------------------
# Reading the raw Defender CSV
# ---------------------------------------------------------------------------

def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_raw_csv(path: str | Path) -> Table:
    """Read the Defender export. The header row is the first row with a
    UserMail column, so any title lines above it are skipped."""
    text = _decode(Path(path).read_bytes())
    rows = list(csv.reader(io.StringIO(text)))
    for i, r in enumerate(rows):
        if any(c.strip().lower() == COL_MAIL.lower() for c in r):
            headers = [c.strip() for c in r]
            data = [row for row in rows[i + 1:] if any(c.strip() for c in row)]
            return Table(headers, data)
    raise ValueError(f"No '{COL_MAIL}' header row found in {Path(path).name} — "
                     "is this the Attack Simulation user report export?")


# ---------------------------------------------------------------------------
# Phish technique + per-user classification
# ---------------------------------------------------------------------------

def _timestamp_columns(headers: list[str]) -> list[str]:
    return [h for h in headers if "_timestamp" in h.lower()]


def unknown_action_columns(headers: list[str]) -> list[str]:
    """*_TimeStamp columns that are neither a known fail action nor a known
    message/notification event — the tab asks about these."""
    return [h for h in _timestamp_columns(headers)
            if not h.split("_", 1)[0].startswith(_FAIL_PREFIXES + _NON_INTERACTION_PREFIXES)]


def interaction_columns(headers: list[str], unknown_fail: set[str] | None = None) -> list[str]:
    """Columns whose timestamp makes a user a Fail.

    unknown_fail: which unknown action columns count as fail (the rest are
    ignored). None = not asked -> every unknown column counts as fail."""
    out = []
    for h in _timestamp_columns(headers):
        prefix = h.split("_", 1)[0]
        if prefix.startswith(_FAIL_PREFIXES):
            out.append(h)
        elif not prefix.startswith(_NON_INTERACTION_PREFIXES):
            if unknown_fail is None or h in unknown_fail:
                out.append(h)
    return out


def detect_phish_type(headers: list[str]) -> str:
    joined = " ".join(headers).lower()
    if "credsupplied" in joined:
        return "Credential Harvest"
    if "attachmentopened" in joined:
        return "Malware Attachment"
    if "emaillinkclicked" in joined:
        return "Drive-by URL"
    return "Unknown"


def _cred_columns(headers: list[str]) -> list[str]:
    return [h for h in headers if h.lower().startswith("credsupplied_timestamp")]


def _interacted(t: Table, row: list, unknown_fail: set[str] | None = None) -> bool:
    if str(t.get(row, COL_COMPROMISED) or "").strip().lower() == "yes":
        return True
    return any(has_value(t.get(row, c)) for c in interaction_columns(t.headers, unknown_fail))


def classify_phase1(t: Table, row: list, unknown_fail: set[str] | None = None) -> str:
    """Sheet for one raw-results row. Order is the fail-first sort order."""
    if has_value(t.get(row, COL_DELIVERY_FAILED)):
        return SHEET_DELIVERY_FAILED
    if _interacted(t, row, unknown_fail) or has_value(t.get(row, COL_REPLIED)):
        return SHEET_FAIL
    if has_value(t.get(row, COL_FORWARDED)):
        return SHEET_FORWARDS
    if has_value(t.get(row, COL_REPORTED)):
        return SHEET_REPORT
    if has_value(t.get(row, COL_DELETED)):
        return SHEET_DELETE
    return SHEET_IGNORE


def fail_reason(t: Table, row: list, unknown_fail: set[str] | None = None) -> str:
    """Why a Fail-sheet row failed, for the email counts.
    Click beats Reply beats Forward (same precedence as the sort). An unknown
    action column answered as fail counts with the clicks / opens."""
    if any(has_value(t.get(row, c)) for c in _cred_columns(t.headers)):
        return "Creds Supplied"
    if _interacted(t, row, unknown_fail):
        return "Click"
    if has_value(t.get(row, COL_REPLIED)):
        return "Reply"
    if has_value(t.get(row, COL_FORWARDED)):
        return "Forward"
    return "Other"


def split_phase1(t: Table, unknown_fail: set[str] | None = None) -> dict[str, Table]:
    sheets = {name: t.empty_like() for name in PHASE1_SHEET_ORDER}
    for row in t.rows:
        sheets[classify_phase1(t, row, unknown_fail)].rows.append(row)
    return sheets


# ---------------------------------------------------------------------------
# File naming
# ---------------------------------------------------------------------------

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "June",
          "July", "Aug", "Sept", "Oct", "Nov", "Dec"]
SIM_TYPES = ["Ev", "Sp"]

_NAME_RE = re.compile(r"^\s*(?P<month>[A-Za-z]+)\s+(?P<year>\d{4})\s+(?P<type>Ev|Sp)\b",
                      re.IGNORECASE)


def parse_report_name(stem: str) -> tuple[str, int, str] | None:
    """'Oct 2026 Ev raw results' -> ('Oct', 2026, 'Ev')."""
    m = _NAME_RE.match(stem)
    if not m:
        return None
    typ = m.group("type").capitalize()
    return m.group("month"), int(m.group("year")), typ


def report_prefix(month: str, year: int | str, sim_type: str) -> str:
    return f"{month.strip()} {year} {sim_type.strip()}"


def phase1_paths(folder: str | Path, prefix: str) -> dict[str, Path]:
    folder = Path(folder)
    return {
        "raw_csv": folder / f"{prefix} raw results.csv",
        "xlsx": folder / f"{prefix} Phish Results.xlsx",
        "forwards_csv": folder / f"{prefix} forwards.csv",
    }


# ---------------------------------------------------------------------------
# Excel writing / reading
# ---------------------------------------------------------------------------

_DT_FORMATS = (
    "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p",
    "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
)
_INT_RE = re.compile(r"^-?\d+$")
_DT_NUMBER_FORMAT = "m/d/yyyy h:mm"


def to_cell_value(v):
    """CSV text -> what Excel would have made of it (dates, whole numbers)."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if s in _EMPTY:
        return s if s else None
    if _INT_RE.match(s) and len(s) < 12:
        return int(s)
    iso = s.replace("Z", "").replace("T", " ") if re.match(r"^\d{4}-\d{2}-\d{2}", s) else None
    if iso:
        try:
            dt = datetime.fromisoformat(iso)
            return dt.replace(tzinfo=None, microsecond=0)
        except ValueError:
            pass
    for fmt in _DT_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return v


def write_workbook(path: str | Path, sheets: dict[str, Table], order: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name in order:
        t = sheets[name]
        ws = wb.create_sheet(name)
        ws.append(list(t.headers))
        for c in ws[1]:
            c.font = bold
        for row in t.rows:
            vals = [to_cell_value(v) for v in row]
            ws.append(vals)
            for c, v in zip(ws[ws.max_row], vals):
                if isinstance(v, datetime):
                    c.number_format = _DT_NUMBER_FORMAT
        ws.freeze_panes = "A2"
        for i, h in enumerate(t.headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(len(h or "") + 2, 30))
    wb.save(path)


def read_workbook(path: str | Path) -> dict[str, Table]:
    """Every sheet as a Table (header = first non-empty row). Keeps tab order."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out: dict[str, Table] = {}
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
            if not rows:
                out[ws.title] = Table([])
                continue
            headers = ["" if v is None else str(v).strip() for v in rows[0]]
            while headers and not headers[-1]:
                headers.pop()
            data = [r[:len(headers)] for r in rows[1:]]
            out[ws.title] = Table(headers, data)
    finally:
        wb.close()
    return out


def workbook_headers(sheets: dict[str, Table]) -> list[str]:
    """Column headers of a report workbook (all its sheets share them)."""
    return next((t.headers for t in sheets.values() if t.headers), [])


def write_forwards_csv(path: str | Path, forwards: Table) -> None:
    idx = [forwards.col(c) for c in FORWARDS_CSV_COLUMNS]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(FORWARDS_CSV_COLUMNS)
        for row in forwards.rows:
            w.writerow(["" if i is None or i >= len(row) or row[i] is None else row[i]
                        for i in idx])


# ---------------------------------------------------------------------------
# Phase 1
# ---------------------------------------------------------------------------

@dataclass
class Phase1Result:
    xlsx: Path
    forwards_csv: Path
    raw_csv: Path
    phish_type: str
    counts: dict[str, int]
    warnings: list[str]


def run_phase1(raw_csv: str | Path, prefix: str, *, rename_raw: bool = True,
               unknown_fail: set[str] | None = None) -> Phase1Result:
    """unknown_fail: unknown action columns (see unknown_action_columns) that
    count as fail; the others are ignored. None = all count as fail."""
    raw_csv = Path(raw_csv)
    t = read_raw_csv(raw_csv)
    warnings: list[str] = []
    missing = [c for c in (COL_MAIL, COL_DELIVERY_FAILED, COL_REPLIED, COL_FORWARDED,
                           COL_REPORTED, COL_DELETED) if t.col(c) is None]
    if missing:
        warnings.append("Missing expected column(s): " + ", ".join(missing)
                        + " — users may land on the wrong sheet. Check the export.")
    phish_type = detect_phish_type(t.headers)
    if phish_type == "Unknown":
        warnings.append("Couldn't tell the phish technique from the columns.")
    unknown = unknown_action_columns(t.headers)
    if unknown and unknown_fail is None:
        warnings.append("Unrecognised action column(s) counted as fail: " + ", ".join(unknown))

    paths = phase1_paths(raw_csv.parent, prefix)
    sheets = split_phase1(t, unknown_fail)
    write_workbook(paths["xlsx"], sheets, PHASE1_SHEET_ORDER)
    write_forwards_csv(paths["forwards_csv"], sheets[SHEET_FORWARDS])

    final_raw = raw_csv
    if rename_raw and raw_csv.resolve() != paths["raw_csv"].resolve():
        if paths["raw_csv"].exists():
            warnings.append(f"Didn't rename the download: {paths['raw_csv'].name} already exists.")
        else:
            raw_csv.rename(paths["raw_csv"])
            final_raw = paths["raw_csv"]

    counts = {name: len(sheets[name].rows) for name in PHASE1_SHEET_ORDER}
    return Phase1Result(paths["xlsx"], paths["forwards_csv"], final_raw,
                        phish_type, counts, warnings)


# ---------------------------------------------------------------------------
# Phase 2 — forward review results
# ---------------------------------------------------------------------------

# Decision for one forwarded user: True = sent to the approved backup
# reporter mailbox (pass), False = didn't (fail), None = not decided yet.

_YES_WORDS = {"yes", "y", "true", "reported", "approved", "pass", "passed", "correct", "ok"}
_NO_WORDS = {"no", "n", "false", "fail", "failed", "incorrect", "wrong"}
_NOT_REPORTED_RE = re.compile(r"\b(not|didn'?t|did not|never)\s+(report|reported|forward)",
                              re.IGNORECASE)


def yes_no(value) -> bool | None:
    """Interpret one cell / phrase as a yes/no, or None if it isn't one."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower().rstrip(".!")
    if not s or s in _EMPTY:
        return None
    if _NOT_REPORTED_RE.search(s):
        return False
    if s in _YES_WORDS:
        return True
    if s in _NO_WORDS:
        return False
    return None


def _decide(values: list[bool | None]) -> bool | None:
    found = {v for v in values if v is not None}
    return found.pop() if len(found) == 1 else None


@dataclass
class ForwardParse:
    decisions: dict[str, bool]          # email -> decision
    ambiguous: list[str]                # emails with conflicting / no yes-no
    unknown_emails: list[str]           # emails in the results that weren't forwarders


def parse_forward_rows(rows: list[list], forward_emails: set[str],
                       listed_means: bool | None = None) -> ForwardParse:
    """Rows from whatever the reviewer sent back (csv/xlsx, any extra columns
    in any position). A row is matched by any cell holding a forwarder's
    email; its decision comes from the other cells that read as yes/no.

    Header row, if present, is used to ignore the original columns we sent
    (so a "--" or date in Phishing Reported On is never read as an answer)."""
    skip: set[int] = set()
    for r in rows[:5]:
        cells = [str(c).strip().lower() if c is not None else "" for c in r]
        if COL_MAIL.lower() in cells:
            skip = {i for i, c in enumerate(cells)
                    if c in {x.lower() for x in FORWARDS_CSV_COLUMNS}}
            break

    decisions: dict[str, bool] = {}
    ambiguous: list[str] = []
    unknown: list[str] = []
    for r in rows:
        emails = []
        for c in r:
            if c is not None:
                emails += [e.lower() for e in _EMAIL_RE.findall(str(c))]
        if not emails:
            continue
        ours = [e for e in emails if e in forward_emails]
        if not ours:
            unknown.extend(e for e in emails if e not in unknown)
            continue
        email = ours[0]
        answers = [yes_no(c) for i, c in enumerate(r)
                   if i not in skip and c is not None and not _EMAIL_RE.search(str(c))]
        d = _decide(answers)
        if d is None and not any(a is not None for a in answers):
            d = listed_means
        if d is None:
            if email not in ambiguous:
                ambiguous.append(email)
        else:
            decisions[email] = d
    return ForwardParse(decisions, [e for e in ambiguous if e not in decisions], unknown)


def parse_forward_text(text: str, forward_emails: set[str],
                       listed_means: bool | None = None) -> ForwardParse:
    """Pasted email body / list. One user per line; Yes/No words on the same
    line decide it. A line with an email and no yes/no uses listed_means."""
    decisions: dict[str, bool] = {}
    ambiguous: list[str] = []
    unknown: list[str] = []
    for line in text.splitlines():
        emails = [e.lower() for e in _EMAIL_RE.findall(line)]
        if not emails:
            continue
        ours = [e for e in emails if e in forward_emails]
        if not ours:
            unknown.extend(e for e in emails if e not in unknown)
            continue
        rest = _EMAIL_RE.sub(" ", line)
        if _NOT_REPORTED_RE.search(rest):
            answers = [False]
        else:
            words = re.findall(r"[A-Za-z']+", rest.lower())
            # a single "y"/"n" only counts when it's the whole answer, not an initial
            answers = [yes_no(w) for w in words if len(w) > 1 or len(words) == 1]
            answers = [a for a in answers if a is not None]
        d = _decide(answers)
        if d is None and not answers:
            d = listed_means
        for email in ours:
            if d is None:
                if email not in ambiguous:
                    ambiguous.append(email)
            else:
                decisions[email] = d
    return ForwardParse(decisions, [e for e in ambiguous if e not in decisions], unknown)


def read_rows_any(path: str | Path) -> list[list]:
    """Rows of a .csv or every sheet of an .xlsx, as plain lists."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=True, read_only=True)
        try:
            return [list(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    return list(csv.reader(io.StringIO(_decode(p.read_bytes()))))


# ---------------------------------------------------------------------------
# Phase 2 — previous simulation history
# ---------------------------------------------------------------------------

PREV_DELIVERY_FAILED = -1


@dataclass
class History:
    streak: dict[str, int]              # email -> prev fail streak (0 = passed)
    warnings: list[str]


def read_history(prev_sheets: dict[str, Table]) -> History:
    """email -> how many times in a row they'd failed as of the previous sim
    (0 if they passed it, PREV_DELIVERY_FAILED if it never reached them)."""
    streak: dict[str, int] = {}
    warnings: list[str] = []
    saw_fail_x = any(_FAIL_X_RE.match(n) for n in prev_sheets)
    for name, t in prev_sheets.items():
        lname = name.strip().lower()
        m = _FAIL_X_RE.match(name)
        if m:
            n = int(m.group(1))
        elif lname.startswith("pass"):
            n = 0
        elif "delivery" in lname and "fail" in lname:
            n = PREV_DELIVERY_FAILED
        elif lname == "fail":
            if saw_fail_x:
                continue        # combined sheet alongside the split ones
            n = 1
            warnings.append("Previous workbook has a plain 'Fail' sheet and no 'Fail x N' "
                            "sheets — treated everyone on it as Fail x 1.")
        else:
            if t.rows:
                warnings.append(f"Ignored sheet '{name}' in the previous workbook.")
            continue
        for row in t.rows:
            e = _norm_email(t.get(row, COL_MAIL))
            if not e:
                continue
            # if someone is on more than one sheet, the worst streak wins
            streak[e] = max(streak.get(e, n), n)
    return History(streak, warnings)


# ---------------------------------------------------------------------------
# Phase 2 — finishing the workbook + summary
# ---------------------------------------------------------------------------

def _pct(n: int, total: int) -> int:
    if not total:
        return 0
    return int((Decimal(n) * 100 / Decimal(total)).quantize(Decimal(1), rounding=ROUND_HALF_UP))


@dataclass
class Summary:
    phish_type: str
    sent: int
    report: int
    delete: int
    ignore: int
    reply: int
    forward: int
    click: int                 # all clicks/opens, creds included ("Attachment Opened" for attachments)
    creds: int
    other: int
    delivery_failed: int
    streaks: dict[int, int]    # Fail x N -> count

    @property
    def passed(self) -> int:
        return self.report + self.delete + self.ignore

    @property
    def failed(self) -> int:
        return self.reply + self.forward + self.click + self.other

    def email_text(self) -> str:
        lines = [
            f"Total Emails Sent: {self.sent}",
            f"\tPassed: {self.passed} ({_pct(self.passed, self.sent)}%)",
            f"\t\tReport: {self.report} ({_pct(self.report, self.sent)}%)",
            f"\t\tDelete: {self.delete}",
            f"\t\tIgnore: {self.ignore}",
            f"\tFail: {self.failed} ({_pct(self.failed, self.sent)}%)",
            f"\t\tReply: {self.reply}",
            f"\t\tForward: {self.forward}",
            f"\t\tClick: {self.click}",
        ]
        if self.phish_type == "Malware Attachment":
            lines[-1] = f"\t\tAttachment Opened: {self.click}"
        if self.phish_type == "Credential Harvest":
            lines += [f"\t\t\tClick-Only: {self.click - self.creds}",
                      f"\t\t\tCreds Supplied: {self.creds}"]
        if self.other:
            lines.append(f"\t\tOther: {self.other}")
        return "\n".join(lines)


def _fail_sheet_names(sheets: dict[str, Table]) -> list[str]:
    return [n for n in sheets if n.strip().lower() == "fail" or _FAIL_X_RE.match(n)]


def summarize(sheets: dict[str, Table], unknown_fail: set[str] | None = None) -> Summary:
    """Counts for the email, from a finished (or phase-1) workbook."""
    def n(name):
        return len(sheets[name].rows) if name in sheets else 0

    any_t = next((t for t in sheets.values() if t.headers), Table([]))
    phish_type = detect_phish_type(any_t.headers)
    reasons = {"Creds Supplied": 0, "Click": 0, "Reply": 0, "Forward": 0, "Other": 0}
    streaks: dict[int, int] = {}
    for name in _fail_sheet_names(sheets):
        t = sheets[name]
        m = _FAIL_X_RE.match(name)
        if m and t.rows:
            streaks[int(m.group(1))] = len(t.rows)
        for row in t.rows:
            reasons[fail_reason(t, row, unknown_fail)] += 1
    pending_forwards = n(SHEET_FORWARDS)
    s = Summary(
        phish_type=phish_type,
        sent=0,
        report=n(SHEET_REPORT), delete=n(SHEET_DELETE), ignore=n(SHEET_IGNORE),
        reply=reasons["Reply"], forward=reasons["Forward"],
        click=reasons["Click"] + reasons["Creds Supplied"], creds=reasons["Creds Supplied"],
        other=reasons["Other"],
        delivery_failed=n(SHEET_DELIVERY_FAILED),
        streaks=dict(sorted(streaks.items())),
    )
    s.sent = s.passed + s.failed + pending_forwards
    return s


@dataclass
class Phase2Result:
    path: Path
    summary: Summary
    warnings: list[str]
    moved_to_report: int
    moved_to_fail: int


def finish_workbook(phase1_xlsx: str | Path, decisions: list[bool],
                    prev_xlsx: str | Path | None, *,
                    out_path: str | Path | None = None,
                    unknown_fail: set[str] | None = None) -> Phase2Result:
    """Apply forward decisions + failure streaks and write the final workbook
    (overwrites the phase-1 workbook unless out_path is given).

    decisions: one per row of the forwards sheet, in sheet order (see
    forwarded_users) — True = forwarded to the approved backup reporter mailbox."""
    phase1_xlsx = Path(phase1_xlsx)
    out_path = Path(out_path) if out_path else phase1_xlsx
    sheets = read_workbook(phase1_xlsx)
    warnings: list[str] = []

    if SHEET_FAIL not in sheets and not any(_FAIL_X_RE.match(n) for n in sheets):
        raise ValueError(f"{phase1_xlsx.name} has no '{SHEET_FAIL}' sheet — pick the phase-1 workbook.")
    if any(_FAIL_X_RE.match(n) for n in sheets):
        raise ValueError(f"{phase1_xlsx.name} already has 'Fail x N' sheets — it looks finished. "
                         "Use 'Recalculate summary' to get the numbers again.")

    template = next(t for t in sheets.values() if t.headers)
    for name in PHASE1_SHEET_ORDER:
        sheets.setdefault(name, template.empty_like())

    # 1. forwarded users -> Pass - Report / Fail
    fwd = sheets[SHEET_FORWARDS]
    moved_r = moved_f = 0
    if len(decisions) != len(fwd.rows) or any(d is None for d in decisions):
        raise ValueError("Every forwarded user needs a result before finishing "
                         f"({len(fwd.rows)} on the forwards sheet).")
    for row, ok in zip(fwd.rows, decisions):
        target = SHEET_REPORT if ok else SHEET_FAIL
        sheets[target].rows.append(_reorder(row, fwd, sheets[target]))
        if ok:
            moved_r += 1
        else:
            moved_f += 1
    del sheets[SHEET_FORWARDS]

    # 2. split Fail by streak
    if prev_xlsx:
        prev_path = Path(prev_xlsx)
        if prev_path.resolve() == phase1_xlsx.resolve():
            raise ValueError("The previous simulation's workbook is the same file as this one.")
        hist = read_history(read_workbook(prev_path))
        warnings += hist.warnings
    else:
        hist = History({}, [])
        warnings.append("No previous simulation workbook — every failure is Fail x 1.")

    fail = sheets.pop(SHEET_FAIL)
    by_streak: dict[int, Table] = {}
    no_email = []
    look_back = []
    for row in fail.rows:
        e = _norm_email(fail.get(row, COL_MAIL))
        prev = hist.streak.get(e, 0) if e else 0
        if not e:
            no_email.append(row)
        if prev == PREV_DELIVERY_FAILED:
            look_back.append(fail.get(row, COL_NAME) or e)
            prev = 0
        by_streak.setdefault(prev + 1, fail.empty_like()).rows.append(row)
    if no_email:
        warnings.append(f"{len(no_email)} failed user(s) have no UserMail — couldn't check "
                        "their history, put on Fail x 1.")
    if look_back:
        warnings.append("Previous sim didn't reach these users (Delivery Failed) — put on "
                        "Fail x 1; check the report before that one: " + "; ".join(map(str, look_back)))

    top = max(by_streak) if by_streak else 1
    order = [SHEET_REPORT, SHEET_DELETE, SHEET_IGNORE]
    for n in range(1, top + 1):
        sheets[fail_x_sheet(n)] = by_streak.get(n, fail.empty_like())
        order.append(fail_x_sheet(n))
    order.append(SHEET_DELIVERY_FAILED)
    # keep any extra sheets she added by hand, at the end
    order += [n for n in sheets if n not in order]

    write_workbook(out_path, sheets, order)
    summary = summarize(sheets, unknown_fail)
    if summary.other:
        warnings.append(f"{summary.other} user(s) on a Fail sheet with no click, reply or forward "
                        "recorded — counted as 'Other'. Check them.")
    return Phase2Result(out_path, summary, warnings, moved_r, moved_f)


def _reorder(row: list, src: Table, dst: Table) -> list:
    """Row from src laid out in dst's column order (sheets normally match)."""
    if src.headers == dst.headers:
        return row
    return [src.get(row, h) for h in dst.headers]


def forwarded_users(phase1_xlsx: str | Path) -> tuple[Table, list[dict]]:
    """The forwards sheet, plus a display row per forwarded user."""
    sheets = read_workbook(phase1_xlsx)
    if SHEET_FORWARDS not in sheets:
        if any(_FAIL_X_RE.match(n) for n in sheets):
            raise ValueError(f"{Path(phase1_xlsx).name} looks finished already (has Fail x N sheets).")
        raise ValueError(f"{Path(phase1_xlsx).name} has no '{SHEET_FORWARDS}' sheet.")
    t = sheets[SHEET_FORWARDS]
    users = [{
        "name": t.get(r, COL_NAME) or "",
        "email": _norm_email(t.get(r, COL_MAIL)),
        "forwarded": t.get(r, COL_FORWARDED),
        "reported": t.get(r, COL_REPORTED),
    } for r in t.rows]
    return t, users
